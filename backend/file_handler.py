"""
file_handler.py

Universal Multimodal Content Transformation Pipeline:
  Stage 1: Ingests & extracts text from any input file (.mp4, .txt, .pdf, .pptx, audio, image)
           with robust OpenCV/FFmpeg fallback for audio-in-MP4 containers.
  Stage 2: LLM transforms source content into selected named deliverables using
           strictly grounded synthesis (banning fabricated checklists and unmentioned steps).
  Stage 3: Physical container builders (.txt, .pdf, .pptx, .mp3) with clean typography.
"""

import os
import re
import json
import time
import base64
import argparse
import subprocess
import cv2
from pypdf import PdfReader
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from groq import Groq
from gtts import gTTS

# ---------------------------------------------------------------------------
# 1. CLIENT SETUP & MODELS
# ---------------------------------------------------------------------------

GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
if not GROQ_API_KEY:
    raise RuntimeError(
        "GROQ_API_KEY environment variable is not set. "
        "PowerShell: $env:GROQ_API_KEY='your_key_here'"
    )

groq_client = Groq(api_key=GROQ_API_KEY)

TEXT_MODEL = "openai/gpt-oss-120b"
FALLBACK_TEXT_MODEL = "openai/gpt-oss-20b"
VISION_MODEL = "qwen/qwen3.8-27b"
AUDIO_MODEL = "whisper-large-v3-turbo"

MANUAL_DESCRIPTION = ""

TEXT_EXTENSIONS = {".txt", ".md", ".log"}
DOC_EXTENSIONS = {".pdf"}
PPT_EXTENSIONS = {".pptx", ".ppt"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm"}

# ---------------------------------------------------------------------------
# 2. LOW-LEVEL HELPERS & TEXT CLEANERS
# ---------------------------------------------------------------------------

def safe_truncate(text: str, max_chars: int = 6000) -> str:
    if len(text) > max_chars:
        print(f"[Warning] Truncating source text to {max_chars} chars to protect quota...")
        return text[:max_chars] + "\n\n...[Truncated]..."
    return text


def clean_plain_text(raw_text: str) -> str:
    """Strips outer codeblocks and clean asterisks while preserving structured text."""
    text = raw_text.strip()
    text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
    text = re.sub(r"\n?```$", "", text)
    text = re.sub(r"\*{1,3}(.*?)\*{1,3}", r"\1", text)

    lines = []
    for line in text.split("\n"):
        clean_l = line.strip()
        if clean_l.startswith("|") and clean_l.endswith("|"):
            cells = [c.strip() for c in clean_l.strip("|").split("|")]
            lines.append("  • " + " — ".join(cells))
        elif re.match(r"^[\-\|\:\s]+$", clean_l):
            continue
        else:
            lines.append(line)
    return "\n".join(lines).strip()


def _clean_json_output(raw_text: str) -> str:
    text = raw_text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text)
        text = text.strip()
    match = re.search(r"(\[.*\]|\{.*\})", text, re.DOTALL)
    return match.group(1).strip() if match else text


def sample_video_keyframes(video_path: str, interval_sec: int = 5, max_frames: int = 3):
    cap = cv2.VideoCapture(video_path)
    fps = cap.get(cv2.CAP_PROP_FPS)
    frames_b64 = []

    if fps == 0:
        cap.release()
        return frames_b64

    frame_interval = int(fps * interval_sec)
    frame_count = 0

    while cap.isOpened() and len(frames_b64) < max_frames:
        ret, frame = cap.read()
        if not ret:
            break
        if frame_count % frame_interval == 0:
            resized = cv2.resize(frame, (640, 360))
            _, buffer = cv2.imencode(".jpg", resized)
            frames_b64.append(base64.b64encode(buffer).decode("utf-8"))
        frame_count += 1

    cap.release()
    return frames_b64


# ---------------------------------------------------------------------------
# STAGE 1: SOURCE INGESTION
# ---------------------------------------------------------------------------

def extract_to_text(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    print(f"[Stage 1] Ingesting & extracting text from: {file_path} ({ext})")

    if ext in TEXT_EXTENSIONS:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    elif ext in DOC_EXTENSIONS:
        reader = PdfReader(file_path)
        pages = [p.extract_text() for p in reader.pages if p.extract_text()]
        return "\n\n".join(pages) if pages else "Notice: PDF parsed, no selectable text found."

    elif ext in PPT_EXTENSIONS:
        prs = Presentation(file_path)
        blocks = []
        for idx, slide in enumerate(prs.slides, start=1):
            lines = [shape.text_frame.text.strip() for shape in slide.shapes
                     if shape.has_text_frame and shape.text_frame.text.strip()]
            blocks.append(f"--- Slide {idx} ---\n" + "\n".join(lines))
        return "\n\n".join(blocks)

    elif ext in AUDIO_EXTENSIONS:
        with open(file_path, "rb") as f:
            return groq_client.audio.transcriptions.create(
                file=(os.path.basename(file_path), f.read()),
                model=AUDIO_MODEL,
                response_format="json"
            ).text

    elif ext in IMAGE_EXTENSIONS:
        with open(file_path, "rb") as img:
            b64 = base64.b64encode(img.read()).decode("utf-8")
        res = groq_client.chat.completions.create(
            model=VISION_MODEL,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": "Provide an exhaustive, accurate textual extraction of everything in this visual."},
                    {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}
                ]
            }],
            max_tokens=700
        )
        return res.choices[0].message.content.strip()

    elif ext in VIDEO_EXTENSIONS:
        cap = cv2.VideoCapture(file_path)
        frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        cap.release()

        if frame_count <= 0:
            print("[Stage 1 Notice] Media container has 0 video frames. Transcribing directly as AUDIO...")
            with open(file_path, "rb") as f:
                return groq_client.audio.transcriptions.create(
                    file=(os.path.basename(file_path), f.read()),
                    model=AUDIO_MODEL,
                    response_format="json"
                ).text

        transcript = ""
        temp_audio = "temp_video_audio.mp3"
        try:
            import imageio_ffmpeg
            ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
            cmd = [ffmpeg_exe, "-y", "-i", file_path, "-vn", "-acodec", "libmp3lame", "-q:a", "4", temp_audio]
            subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
            if os.path.exists(temp_audio) and os.path.getsize(temp_audio) > 100:
                with open(temp_audio, "rb") as f:
                    transcript = groq_client.audio.transcriptions.create(
                        file=(temp_audio, f.read()),
                        model=AUDIO_MODEL,
                        response_format="json"
                    ).text
        except Exception as audio_err:
            print(f"[Warning] Audio track extraction notice: {audio_err}")
        finally:
            if os.path.exists(temp_audio):
                try:
                    os.remove(temp_audio)
                except Exception:
                    pass

        visual_desc = ""
        try:
            frames = sample_video_keyframes(file_path, interval_sec=5, max_frames=3)
            if frames:
                content = [{"type": "text", "text": "Describe the core actions, slides, or scene transitions in these frames."}]
                for b64 in frames:
                    content.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})
                res = groq_client.chat.completions.create(model=VISION_MODEL, messages=[{"role": "user", "content": content}])
                visual_desc = res.choices[0].message.content.strip()
        except Exception as e:
            print(f"[Warning] Keyframe visual extraction notice: {e}")

        combined = ""
        if transcript.strip():
            combined += f"SPOKEN DIALOGUE TRANSCRIPT:\n{transcript}\n\n"
        if visual_desc.strip():
            combined += f"VISUAL SCENE CONTEXT:\n{visual_desc}"

        return combined.strip() if combined.strip() else "Notice: Media file parsed, no speech or frames identified."

    else:
        raise ValueError(f"Unsupported format: {ext}")


# ---------------------------------------------------------------------------
# STAGE 2: STRICTLY GROUNDED DELIVERABLE PROMPT CONTRACTS
# ---------------------------------------------------------------------------

OUTPUT_TYPES = {
    "plain_summary": {
        "container": "text",
        "template": (
            "Produce an executive text summary synthesized directly from the provided source.\n\n"
            "STRICT RULES:\n"
            "- Ground every sentence in the source. Do not extrapolate external theories or frameworks.\n"
            "- No meta-commentary ('The author states...', 'This text discusses...'). Begin directly.\n"
            "- Do not use asterisks (**) or markdown table borders.\n\n"
            "REQUIRED STRUCTURE:\n\n"
            "CORE MESSAGE & PREMISE\n"
            "State the central argument and fundamental reality in 2-3 clear, powerful sentences.\n\n"
            "KEY PRINCIPLES & FINDINGS\n"
            "- Direct insight, critical distinction, or warning explicitly stated in the source.\n"
            "- Secondary operational observation or dynamic detailed in the text.\n"
            "- Core mechanism or relationship highlighted by the author.\n\n"
            "STRATEGIC TAKEAWAYS\n"
            "- Conclusive guideline or direct takeaway derived solely from the text."
        ),
    },
    "audio_briefing": {
        "container": "audio",
        "template": (
            "Write a natural, conversational executive audio briefing script summarizing the core message.\n\n"
            "STRICT RULES:\n"
            "- Write in first-person or direct professional narrator tone meant to be read aloud.\n"
            "- Ground every sentence in the provided source without external speculation.\n"
            "- Zero meta-commentary (never say 'In this audio', 'The author says', or 'Thank you for listening').\n"
            "- Zero asterisks (**), zero markdown symbols, zero bullet points.\n"
            "- 2 to 3 flowing, spoken paragraphs designed for clear vocal delivery."
        ),
    },
    "linkedin_post": {
        "container": "text",
        "template": (
            "Write a high-impact, professional LinkedIn post based directly on the core lesson of the content.\n\n"
            "STRICT RULES:\n"
            "- Line 1 MUST be a bold hook capturing the core insight. No greetings or introductory filler.\n"
            "- Zero meta-commentary (never say 'In a recent speech' or 'Here is a lesson').\n"
            "- Deliver 3 to 4 concise, punchy paragraphs with clean spacing.\n"
            "- Avoid markdown asterisks (**). Conclude with exactly 3 relevant hashtags at the end."
        ),
    },
    "x_thread": {
        "container": "text",
        "template": (
            "Distill the core premise and vital takeaways into a cohesive 5-tweet thread.\n\n"
            "STRICT RULES:\n"
            "- Return ONLY a valid JSON array of 5 strings: [\"Tweet 1 text\", \"Tweet 2 text\", ...].\n"
            "- Tweet 1 must be a compelling thesis hook. Tweet 5 must be a memorable concluding rule.\n"
            "- Each tweet must be under 260 characters, natural, and completely free of asterisks (**) or meta-commentary."
        ),
    },
    "advisory": {
        "container": "pdf",
        "template": (
            "You are a strategic intelligence consultant. Translate the source material into an authoritative, publication-grade Advisory Document.\n\n"
            "STRICT FIDELITY RULES:\n"
            "- Ground all guidance strictly on observations and distinctions present in the source.\n"
            "- ABSOLUTE PROHIBITION: Do NOT invent corporate audit steps, 30-day plans, SMART KPIs, or daily tracking checklists not stated by the author.\n"
            "- Do NOT output raw bracketed tags like [SECTION: ...]. Use clean Markdown headings.\n\n"
            "REQUIRED STRUCTURE:\n"
            "## ADVISORY: [Clear, Authoritative Title Derived from Source]\n\n"
            "### 1. Executive Context & Foundational Thesis\n"
            "[Detailed narrative establishing the background, core premise, and overarching dynamic.]\n\n"
            "### 2. Core Principles & Analytical Insights\n"
            "- [First primary principle or systemic dynamic directly from source]\n"
            "- [Second primary principle or systemic dynamic directly from source]\n"
            "- [Third primary principle or systemic dynamic directly from source]\n\n"
            "### 3. Critical Vulnerabilities & Failure Modes\n"
            "[Examine what happens when the core rule or constraint is ignored, as explained in the content.]\n\n"
            "### 4. Strategic Directives\n"
            "- [Authoritative strategic principle or rule of action established by the text]\n"
            "- [Authoritative strategic principle or rule of action established by the text]"
        ),
    },
    "executive_summary": {
        "container": "pdf",
        "template": (
            "You are a senior executive briefing specialist. Synthesize the provided content into a formal, decision-ready Executive Summary.\n\n"
            "STRICT RULES:\n"
            "- No developer bracket tags like [SECTION: ...]. Use formal Markdown headings.\n"
            "- Maintain an objective, concise, decision-oriented tone.\n"
            "- Draw strictly from the ingested content without inventing external timelines, tasks, or audits.\n\n"
            "REQUIRED STRUCTURE:\n"
            "## EXECUTIVE SUMMARY: [Clear Title Derived from Source]\n\n"
            "### 1. Strategic Context\n"
            "[A concise narrative paragraph defining the core situational premise and background.]\n\n"
            "### 2. Key Insights & Critical Observations\n"
            "- [Core observation and operational relevance directly from source]\n"
            "- [Secondary observation and operational relevance directly from source]\n"
            "- [Tertiary observation and operational relevance directly from source]\n\n"
            "### 3. Strategic Implication\n"
            "[A focused concluding takeaway summarizing the overarching strategic lesson or direction.]"
        ),
    },
    "infographic": {
        "container": "text",
        "template": (
            "Convert the source content into a structured visual infographic specification.\n\n"
            "STRICT RULES:\n"
            "- Return ONLY valid JSON in this exact structure:\n"
            "{\n"
            '  "headline": "Bold, punchy visual headline",\n'
            '  "key_stats": ["Key statistic, verified metric, or memorable core takeaway phrase from text - NEVER invent fake percentages"],\n'
            '  "sections": [\n'
            '    {"heading": "Panel 1 Title", "text": "Concise panel copy grounded in text."},\n'
            '    {"heading": "Panel 2 Title", "text": "Concise panel copy grounded in text."},\n'
            '    {"heading": "Panel 3 Title", "text": "Concise panel copy grounded in text."}\n'
            "  ],\n"
            '  "layout_notes": "Recommended visual hierarchy, color accents, and icon directions."\n'
            "}"
        ),
    },
    "presentation": {
        "container": "pptx",
        "template": (
            "Convert the core message and analysis of the source text into an executive 5-slide presentation.\n\n"
            "STRICT RULES:\n"
            "- Return ONLY valid JSON with this exact schema:\n"
            "{\n"
            '  "title": "Comprehensive Deck Title",\n'
            '  "slides": [\n'
            '    {"title": "Slide 1 Title", "bullets": ["Point 1", "Point 2", "Point 3"], "notes": "Speaker talking points."},\n'
            '    {"title": "Slide 2 Title", "bullets": ["Point 1", "Point 2", "Point 3"], "notes": "Speaker talking points."},\n'
            '    {"title": "Slide 3 Title", "bullets": ["Point 1", "Point 2", "Point 3"], "notes": "Speaker talking points."},\n'
            '    {"title": "Slide 4 Title", "bullets": ["Point 1", "Point 2", "Point 3"], "notes": "Speaker talking points."},\n'
            '    {"title": "Slide 5 Title", "bullets": ["Point 1", "Point 2", "Point 3"], "notes": "Speaker talking points."}\n'
            "  ]\n"
            "}\n"
            "- Every bullet must be a substantive statement of principle or fact from the text, not meta-narrator commentary."
        ),
    },
    "video_package": {
        "container": "text",
        "template": (
            "Produce a production-grade 60-90 second Video Production Package text document translating the source content.\n\n"
            "STRICT RULES:\n"
            "- Do not add pleasantries ('Here is your script'). Start directly with the title.\n"
            "- Avoid markdown asterisks (**).\n\n"
            "REQUIRED STRUCTURE:\n\n"
            "============================================================\n"
            "VIDEO PRODUCTION SPECIFICATION\n"
            "============================================================\n\n"
            "TITLE: [Catchy, impactful title]\n"
            "ESTIMATED RUNTIME: 60-90 Seconds\n"
            "CORE PREMISE: [1-2 sentences on what this video communicates]\n\n"
            "------------------------------------------------------------\n"
            "1. COMPLETE NARRATION SCRIPT\n"
            "------------------------------------------------------------\n"
            "(Full spoken voiceover script in natural, engaging cadence)\n\n"
            "------------------------------------------------------------\n"
            "2. SCENE-BY-SCENE STORYBOARD\n"
            "------------------------------------------------------------\n"
            "Scene 1 [00:00 - 00:15]:\n"
            "  Visual: [Actionable visual scene description]\n"
            "  On-Screen Text: [Concise text overlay]\n"
            "  Audio Cue: [Tone/music cue]\n\n"
            "Scene 2 [00:15 - 00:35]:\n"
            "  Visual: [Actionable visual scene description]\n"
            "  On-Screen Text: [Concise text overlay]\n"
            "  Audio Cue: [Tone/music cue]\n\n"
            "Scene 3 [00:35 - 00:60]:\n"
            "  Visual: [Actionable visual scene description]\n"
            "  On-Screen Text: [Concise text overlay]\n"
            "  Audio Cue: [Tone/music cue]\n\n"
            "------------------------------------------------------------\n"
            "3. TIMED SUBTITLES\n"
            "------------------------------------------------------------\n"
            "[00:00 - 00:08] Subtitle line 1\n"
            "[00:08 - 00:18] Subtitle line 2\n"
            "[00:18 - 00:30] Subtitle line 3\n\n"
            "------------------------------------------------------------\n"
            "4. B-ROLL & VISUAL ASSET GUIDELINES\n"
            "------------------------------------------------------------\n"
            "- Asset recommendation 1\n"
            "- Asset recommendation 2"
        ),
    },
}

OUTPUT_TYPES["infographics"] = OUTPUT_TYPES["infographic"]


def generate_output(
    source_text: str,
    output_type: str,
    audience: str = "general",
    tone: str = "professional",
    language: str = "English",
    detail_level: str = "moderate",
    objective: str = None,
    style: str = None,
    custom_instruction: str = None,
) -> str:
    resolved_type = "infographic" if output_type == "infographics" else output_type
    if resolved_type not in OUTPUT_TYPES:
        raise ValueError(f"Unknown output_type '{output_type}'. Valid: {list(OUTPUT_TYPES.keys())}")

    spec = OUTPUT_TYPES[resolved_type]
    container = spec["container"]

    params_block = (
        f"Target Audience: {audience}\n"
        f"Tone: {tone}\n"
        f"Language: {language}\n"
        f"Detail Level: {detail_level}\n"
    )
    if objective:
        params_block += f"Communication Objective: {objective}\n"
    if style:
        params_block += f"Content Style: {style}\n"

    extra = ""
    if custom_instruction and custom_instruction.strip():
        extra = f"\nOPERATOR STEERING & FOCUS DIRECTIVES (PRIORITIZE THIS):\n{custom_instruction.strip()}\n"

    system_instruction = (
        "You are an expert executive content transformation engine. "
        "STRICT GROUNDING DIRECTIVE: Rely exclusively on the provided source content. "
        "Never invent outside frameworks, personal audits, checklists, or procedural steps not in the source text. "
        "Never output conversational pleasantries, introductory remarks, or markdown asterisks (**)."
    )

    prompt_body = (
        f"CONTENT TYPE CONTRACT:\n{spec['template']}\n\n"
        f"GENERATION PARAMETERS:\n{params_block}"
        f"{extra}\n"
        f"SOURCE CONTENT:\n{safe_truncate(source_text)}"
    )

    output = ""
    try:
        res = groq_client.chat.completions.create(
            model=TEXT_MODEL,
            messages=[
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": prompt_body}
            ],
            temperature=0.2,
            max_tokens=1500
        )
        output = res.choices[0].message.content.strip()
    except Exception as primary_err:
        print(f"[Warning] Primary model '{TEXT_MODEL}' failed: {primary_err}")
        print(f"[Notice] Retrying with fallback model '{FALLBACK_TEXT_MODEL}'...")
        try:
            res = groq_client.chat.completions.create(
                model=FALLBACK_TEXT_MODEL,
                messages=[
                    {"role": "system", "content": system_instruction},
                    {"role": "user", "content": prompt_body}
                ],
                temperature=0.2,
                max_tokens=1500
            )
            output = res.choices[0].message.content.strip()
        except Exception as fallback_err:
            raise RuntimeError(f"Both text models failed. Last error: {fallback_err}")

    if container == "pptx" or resolved_type in ("x_thread", "infographic"):
        return _clean_json_output(output)

    return output


# ---------------------------------------------------------------------------
# STAGE 3: EXECUTIVE CONTAINER BUILDERS
# ---------------------------------------------------------------------------

def export_to_txt(text_content: str, output_path: str) -> str:
    cleaned = clean_plain_text(text_content)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(cleaned)
    abs_path = os.path.abspath(output_path)
    print(f"[Export] Clean Text deliverable saved: {abs_path}")
    return abs_path


def export_to_mp3(text_content: str, output_path: str, language: str = "English") -> str:
    """Converts clean synthesis text into an MP3 speech briefing using gTTS."""
    clean_text = clean_plain_text(text_content)
    lang_map = {
        "English": "en",
        "Tamil": "ta",
        "Hindi": "hi",
        "Malayalam": "ml",
        "Telugu": "te"
    }
    tts_lang = lang_map.get(language, "en")
    tts = gTTS(text=clean_text, lang=tts_lang, slow=False)
    tts.save(output_path)
    abs_path = os.path.abspath(output_path)
    print(f"[Export] Voice Audio deliverable saved: {abs_path}")
    return abs_path


def format_infographic_text(json_text: str) -> str:
    try:
        data = json.loads(json_text)
    except Exception:
        clean = _clean_json_output(json_text)
        try:
            data = json.loads(clean)
        except Exception:
            return json_text

    lines = []
    lines.append("=" * 60)
    lines.append("INFOGRAPHIC CREATIVE SPECIFICATION & COPY")
    lines.append("=" * 60)
    lines.append(f"\nHEADLINE: {data.get('headline', '')}")
    lines.append("")
    lines.append("-" * 60)
    lines.append("KEY STATISTICS & CORE TAKEAWAYS")
    lines.append("-" * 60)
    for stat in data.get("key_stats", []):
        lines.append(f"  • {stat}")
    lines.append("")
    lines.append("-" * 60)
    lines.append("CONTENT SECTIONS & PANELS")
    lines.append("-" * 60)
    for i, sec in enumerate(data.get("sections", []), start=1):
        lines.append(f"\nPanel {i}: {sec.get('heading', '')}")
        lines.append(f"Copy: {sec.get('text', '')}")
    lines.append("")
    lines.append("-" * 60)
    lines.append("LAYOUT & VISUAL DESIGN RECOMMENDATIONS")
    lines.append("-" * 60)
    lines.append(f"Architecture: {data.get('layout_notes', '')}")
    return "\n".join(lines).strip()


def format_x_thread_text(json_text: str) -> str:
    try:
        tweets = json.loads(json_text)
    except Exception:
        clean = _clean_json_output(json_text)
        try:
            tweets = json.loads(clean)
        except Exception:
            raw_lines = [l.strip() for l in json_text.split("\n") if l.strip()]
            return "\n\n".join(raw_lines)

    lines = []
    total = len(tweets)
    for i, tweet in enumerate(tweets, start=1):
        lines.append(f"Tweet {i}/{total}:")
        lines.append(str(tweet).replace("**", ""))
        lines.append("")
    return "\n".join(lines).strip()


def export_to_pdf(text_content: str, output_path: str, title: str = "Executive Briefing") -> str:
    doc = SimpleDocTemplate(
        output_path, pagesize=letter,
        rightMargin=48, leftMargin=48, topMargin=44, bottomMargin=44
    )
    styles = getSampleStyleSheet()
    story = []

    PRIMARY_COLOR = colors.HexColor("#0F2942")
    ACCENT_COLOR = colors.HexColor("#2B6CB0")
    BODY_COLOR = colors.HexColor("#2D3748")
    MUTED_COLOR = colors.HexColor("#718096")

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=PRIMARY_COLOR,
        spaceAfter=4
    )

    section_style = ParagraphStyle(
        'DocSection',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=ACCENT_COLOR,
        spaceBefore=14,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14.5,
        textColor=BODY_COLOR,
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=BODY_COLOR,
        leftIndent=16,
        firstLineIndent=-10,
        spaceAfter=4
    )

    story.append(Paragraph(title.upper(), title_style))
    story.append(Paragraph("Intelligent Content Transformation Pipeline • Strategic Deliverable",
                           ParagraphStyle('Sub', parent=body_style, fontSize=8, textColor=MUTED_COLOR, spaceAfter=8)))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY_COLOR, spaceBefore=2, spaceAfter=12))

    clean_text = clean_plain_text(text_content)
    lines = clean_text.split("\n")

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue

        is_heading = False
        heading_text = ""

        m_md = re.match(r"^#{1,4}\s+(.*)$", line)
        m_sec = re.match(r"^\[SECTION:\s*(.*?)\]$", line, re.IGNORECASE)
        m_num = re.match(r"^(?:SECTION\s*\d*[:\.\-]?\s*|\d+[\.\)]\s+)([A-Z0-9\s\-\:\.\(\)\&]+)$", line)

        if m_md:
            heading_text = m_md.group(1).strip()
            is_heading = True
        elif m_sec:
            heading_text = m_sec.group(1).strip()
            is_heading = True
        elif m_num and len(line) < 70:
            heading_text = line
            is_heading = True

        if is_heading:
            safe_heading = heading_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Spacer(1, 4))
            story.append(Paragraph(safe_heading, section_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceBefore=2, spaceAfter=6))
            continue

        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        if safe_line.startswith("- ") or safe_line.startswith("• "):
            bullet_text = safe_line.lstrip("-• ").strip()
            story.append(Paragraph(f"&bull;  {bullet_text}", bullet_style))
        else:
            story.append(Paragraph(safe_line, body_style))

    doc.build(story)
    abs_path = os.path.abspath(output_path)
    print(f"[Export] Professional PDF saved: {abs_path}")
    return abs_path


def export_to_pptx(json_text: str, output_path: str) -> str:
    try:
        data = json.loads(json_text)
    except Exception:
        clean = _clean_json_output(json_text)
        data = json.loads(clean)

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = data.get("title", "Executive Presentation")

    for slide_data in data.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = slide_data.get("title", "")
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide_data.get("bullets", [])):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(bullet).replace("**", "")
            p.font.size = Pt(17)
        if slide_data.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_data["notes"]

    prs.save(output_path)
    abs_path = os.path.abspath(output_path)
    print(f"[Export] Presentation saved: {abs_path}")
    return abs_path


def build_deliverable(generated_text: str, output_type: str, output_dir: str = ".", language: str = "English"):
    os.makedirs(output_dir, exist_ok=True)
    resolved_type = "infographic" if output_type == "infographics" else output_type
    container = OUTPUT_TYPES[resolved_type]["container"]

    if container == "pdf":
        title = resolved_type.replace("_", " ").title()
        return export_to_pdf(generated_text, os.path.join(output_dir, f"{resolved_type}.pdf"), title=title)
    elif container == "pptx":
        return export_to_pptx(generated_text, os.path.join(output_dir, f"{resolved_type}.pptx"))
    elif container == "audio":
        return export_to_mp3(generated_text, os.path.join(output_dir, f"{resolved_type}.mp3"), language=language)
    elif resolved_type == "infographic":
        formatted = format_infographic_text(generated_text)
        return export_to_txt(formatted, os.path.join(output_dir, f"{resolved_type}.txt"))
    elif resolved_type == "x_thread":
        formatted = format_x_thread_text(generated_text)
        return export_to_txt(formatted, os.path.join(output_dir, f"{resolved_type}.txt"))
    else:  # plain_summary, linkedin_post, video_package
        return export_to_txt(generated_text, os.path.join(output_dir, f"{resolved_type}.txt"))


# ---------------------------------------------------------------------------
# PIPELINE RUNNER
# ---------------------------------------------------------------------------

def run_pipeline(
    file_path: str,
    selected_outputs: list,
    output_dir: str = ".",
    audience: str = "general",
    tone: str = "professional",
    language: str = "English",
    detail_level: str = "moderate",
    objective: